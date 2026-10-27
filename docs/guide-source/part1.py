from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docbuild import INK, MUTED


def build(g):
    d = g.doc
    # ── Title page ────────────────────────────────────────────────────────────
    for _ in range(6):
        d.add_paragraph()
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("Interlibrary Loan (ILL) System"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = INK
    s = d.add_paragraph()
    r = s.add_run("Installation, Configuration and Operations Guide"); r.font.size = Pt(16); r.font.color.rgb = MUTED
    d.add_paragraph()
    g.p("A WordPress plugin, an ILL database and a set of scheduled scripts that together run a regional "
        "interlibrary loan service: libraries find an item in a shared catalog, request it from another library, "
        "and track it until it is returned.", color=MUTED)
    for _ in range(8):
        d.add_paragraph()
    g.p("**Plugin version:** 3.13.0", color=MUTED)
    g.p("**Guide date:** October 2026", color=MUTED)
    g.p("**Audience:** system administrators who install and host the service, the ILL coordinator who configures "
        "and runs it, and library staff who use it every day.", color=MUTED)

    g.page_break()
    g.doc.add_heading("Contents", level=1)
    g.toc()

    # ── 1. Overview ───────────────────────────────────────────────────────────
    g.h1("1. Overview")
    g.p("The ILL System lets a group of libraries (a council, consortium or library system) lend to one another. "
        "It is built so any organization can run it: every name, color, address, rule and list is a setting, "
        "and nothing in the code has to be changed to make it yours.")
    g.h2("1.1 What the system does")
    g.bullets([
        "**Requesting.** Staff search the shared catalog, choose an item and send a request to one of the libraries "
        "that owns it. Optionally, the request moves to a backup library automatically if the first one can't fill it.",
        "**Lending.** The lending library gets an email with Yes / No buttons and answers on the website, "
        "with a due date and shipping method, or a reason it can't fill.",
        "**Tracking.** Both sides follow the item through received, renewed, returned and checked in. "
        "Every step is recorded against the request.",
        "**Automatic follow-up.** Scheduled scripts remind lenders who have not answered, expire requests that "
        "are never answered, remind borrowers before items are due and send weekly statistics.",
        "**ILLiad.** Libraries that run ILLiad receive requests directly in ILLiad, and the status, due date and "
        "returns are copied back automatically.",
        "**Administration.** A library directory, user accounts with approval, statistics, data exports, "
        "and a built-in security suite (login protection, two-factor authentication with security keys, passkeys or an authenticator app, file-change monitoring).",
    ])
    g.h2("1.2 The parts")
    g.table(["Part", "What it is", "Where it lives"], [
        ["WordPress site", "The website staff log in to. Any standard WordPress theme can be used.", "Web server, WordPress database"],
        ["ILL System plugin", ["All ILL pages, settings and emails.", "Installed like any WordPress plugin."], "wp-content/plugins/"],
        ["ILL database", ["The request table, the library directory, the ILLiad address mapping and the fallback table.",
                          "Can be on the same database server as WordPress, or a separate one."], "MySQL / MariaDB"],
        ["Scheduled scripts", ["Eight PHP command-line scripts run by cron for reminders, expiry, ILLiad checks and statistics.",
                               "They do not load WordPress; they read its settings directly from its database."], "A folder outside the web root, e.g. /var/www/<code>_script"],
        ["Catalog search service", "A union catalog / Z39.50 search gateway that lists which libraries own a title, with their holdings.", "Your catalog vendor"],
        ["ILLiad (optional)", "Atlas Systems ILLiad web platform API for lenders that use ILLiad.", "Each lending library"],
        ["Mail relay", "An SMTP service that delivers the system's email.", "Your mail provider"],
    ], [1.35, 3.25, 1.9])
    g.h2("1.3 How a request moves through the system")
    g.steps([
        "A borrower searches the catalog and clicks **Request** on a title. The request page lists every participating "
        "library that owns an available copy (rules decide which libraries are shown; see section 4.8).",
        "The borrower picks a lending library, adds a need-by date and notes, and submits. The system gives the request "
        "an ILL number (the year and a sequence number, e.g. 2026-18120), emails the lender with Yes / No buttons, "
        "and emails the borrower a copy. ILLiad lenders also receive it in ILLiad.",
        "The lender answers **Yes, will fill** (due date and shipping method) or **No, can't fill** (a reason). "
        "The borrower is emailed either way. If the borrower turned on Automatic Fallback, a No moves the request "
        "to the next library on their list.",
        "No answer after 3 business days: the lender gets a reminder. No answer after 5 business days: the request "
        "expires (or, with fallback on, moves to the next library).",
        "The borrower marks the item **Received**. Before it is due they can **Request a Renewal**; the lender "
        "approves (with a new due date) or denies. The borrower gets a reminder 5 days before the due date.",
        "The borrower marks the item **Returned**, and the lender **Checks it back in**. The request is complete.",
    ])
    g.h2("1.4 Who uses it")
    g.table(["Role", "Who", "Can do"], [
        ["Library Staff", "Staff at a member library",
         ["Request items, answer requests sent to their library, renew, mark received / returned / checked in,",
          "view their own library's history and statistics, edit their profile and their library's profile."]],
        ["Library System Staff", "Staff at a library system or regional office",
         "Everything Library Staff can do, plus all requests, accounts, library records and statistics for the libraries in their own system."],
        ["ILL Administrator", "The ILL coordinator",
         ["Every ILL administration page: all libraries, all user accounts, all requests, statistics, exports, ILLiad mapping, "
          "resending requests and the fallback overview.",
          "Cannot change WordPress itself, the ILL settings screens, the database password or the security settings, "
          "and cannot edit or create Administrators."]],
        ["Administrator", "The person who hosts and configures the site", "Everything, including all settings under the ILL System menu in the WordPress dashboard."],
    ], [1.4, 1.6, 3.5])
    g.p("A user's library is identified by its **LOC code**: a short code (often a MARC organization code) stored on the "
        "user's profile and on the library's record. One account can be given extra LOC codes so it can act for several "
        "libraries; the pages then offer a library selector.")

    # ── 2. Requirements ───────────────────────────────────────────────────────
    g.h1("2. Requirements")
    g.h2("2.1 Server and software")
    g.table(["Item", "Requirement", "Notes"], [
        ["Operating system", "Linux (any current distribution)", "The scheduled scripts use cron and the curl command."],
        ["Web server", "Apache 2.4 (recommended) or nginx", "Apache reads the plugin's .htaccess, which blocks its notes and scripts from the web. On nginx, add the equivalent rules (section 3.5)."],
        ["PHP", "8.0 or later; 8.1 or later recommended", "Same version for the web server and the command line (cron)."],
        ["PHP extensions", "mysqli, openssl, sodium, curl, mbstring, xml / simplexml, json", "sodium (or openssl) encrypts the stored ILL database password. openssl also checks security keys and passkeys."],
        ["WordPress", "6.0 or later", "Pretty permalinks (e.g. Post name) recommended."],
        ["Database", "MySQL 5.7+ or MariaDB 10.3+", "WordPress and the ILL data may share a server; use separate database accounts."],
        ["HTTPS", "A valid certificate for the site's address", "Required for logins, two-factor and email links. Security keys and passkeys only work over HTTPS on the site's own address."],
        ["Mail", "An SMTP relay that is allowed to send for your From address", "WordPress sends through an SMTP plugin; the scheduled scripts send through the server's mail program (sendmail / Postfix relay)."],
    ], [1.3, 2.2, 3.0])
    g.h2("2.2 Outbound connections")
    g.p("The server must be able to reach these services over HTTPS. Firewalls and security groups often block outbound "
        "traffic by default, so check each one.")
    g.table(["Service", "Used for", "Needed?"], [
        ["Your catalog search gateway", "Reading holdings for the request page", "Yes"],
        ["Each ILLiad library's web platform", "Sending requests to ILLiad and reading status back", "Only if any library uses ILLiad"],
        ["date.nager.at", "Public holidays for business-day counting (scheduled scripts)", "Optional; can be turned off"],
        ["api.pwnedpasswords.com", "Breached-password check at login", "Only if that security option is on"],
        ["api.wordpress.org", "Checking WordPress core files against the official release (file monitor)", "Only if the file monitor is on"],
        ["Your SMTP relay", "All email", "Yes"],
    ], [2.1, 3.0, 1.4])
    g.h2("2.3 What to have ready")
    g.bullets([
        "Your organization's names (short and full), web address, logo and brand colors.",
        "The email addresses to use: From address, alert/support addresses, and who gets the weekly statistics.",
        "A list of member libraries with LOC code, name, library system, ILL email address and mailing address.",
        "Your library systems and couriers / delivery methods.",
        "The catalog search gateway address from your catalog vendor, and how each catalog's location names map to your libraries.",
        "For ILLiad lenders: their ILLiad web platform address, API key and address-mapping details.",
    ])

    # ── 3. Installation ───────────────────────────────────────────────────────
    g.h1("3. Installation")
    g.p("Install in this order. A test site is strongly recommended before going live; turn on **Email test mode** "
        "(section 4.3) on any test site so no real library receives email.")
    g.h2("3.1 Create the ILL database")
    g.p("Create a database and a user for the ILL data. The plugin needs SELECT, INSERT, UPDATE and DELETE on its tables, "
        "plus CREATE the first time the Automatic Fallback module is turned on (it creates its own table).")
    g.code("""CREATE DATABASE ill_data CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;
CREATE USER 'ill_app'@'%' IDENTIFIED BY 'a-long-random-password';
GRANT SELECT, INSERT, UPDATE, DELETE, CREATE ON ill_data.* TO 'ill_app'@'%';
FLUSH PRIVILEGES;""")
    g.p("You don't need to create the tables by hand. Once the plugin is connected (section 3.4), **ILL System → Database → "
        "Run database check** lists the tables that are missing and **Create the missing tables** creates them, empty, under the "
        "names set on that screen (defaults STATS, Library-Data, ILLIAD-ADD-MAPPING and ILL-ROTA). This is why the user needs "
        "CREATE. The check also shows the exact CREATE TABLE statements, if you would rather run them yourself.")
    g.p("If you are moving from an older ILL system that already has these tables, keep them: the plugin reads the same columns. "
        "Run the database check after connecting; it lists any column that is missing or too small, with the SQL that fixes it. "
        "The full column list is in Appendix A.")
    g.note("Optional delivery columns (`lbsyscourier`, `lbUSPS` and a regional-courier column) and the `periodical_loan` column can be "
           "added to the library table later. The library form only shows the questions whose columns exist, so you can leave out "
           "the ones you don't use. The database check shows the SQL for each.", "tip")

    g.h2("3.2 Install WordPress")
    g.steps([
        "Install WordPress 6.0 or later with HTTPS, in its own database.",
        "Install and activate your theme. Pages built by the plugin appear inside the theme's normal content area, "
        "with an ILL navigation sidebar.",
        "Install an SMTP plugin (for example WP Mail SMTP or FluentSMTP) and send a test email. The plugin sends all of its "
        "own mail through WordPress, so this is what delivers request emails.",
        "Under **Settings → Permalinks**, choose **Post name** and save.",
        "Under **Settings → General**, set the time zone to your local city (for example America/Chicago). "
        "Dates on requests and business-day counting use it.",
    ])
    g.h2("3.3 Install the plugin")
    g.steps([
        "Upload the plugin folder to `wp-content/plugins/` (or use **Plugins → Add New → Upload Plugin** with the zip file).",
        "Activate it under **Plugins**. Activation creates the Library Staff, Library System Staff and ILL Administrator roles.",
        "Go to **Settings → Permalinks** and click **Save** once more. This registers the ILL page addresses (/request, /respond, ...).",
        "A new **ILL System** menu appears in the WordPress dashboard. (Once you set the System Short Name it is labelled with that name.)",
        "If you are replacing older separate plugins that did the same jobs (custom user fields, login redirect, terms of service, "
        "catalog search header, user export, dashboard quick links, user approval or last-login tracking), deactivate and delete them. "
        "Those features are built in now, and running both would apply them twice.",
    ])
    g.h2("3.4 Connect the ILL database")
    g.steps([
        "Open **ILL System → Database**.",
        "Enter the host, user name, database name and password of the ILL database from section 3.1, and save.",
        "If your table names differ from the defaults, change them under **Table Names**.",
        "If the database is on another server, turn on **Require Encrypted Connection** and give the CA certificate path "
        "(see **MySQL TLS / SSL Encryption** on the same page).",
        "Open **ILL System → Dashboard**. It confirms the connection and shows counts.",
        "Back on **ILL System → Database**, select **Run database check** (below the form). On a new install, select **Create the "
        "missing tables**, then run the check again. Fix anything marked Problem before going further (section 4.4).",
    ])
    g.note("The database password is stored encrypted, using a key derived from WordPress's secret keys in wp-config.php. "
           "If you ever change those secret keys, re-enter the ILL database password on this screen.", "warn")
    g.h2("3.5 Protect the plugin folder")
    g.p("The plugin folder contains notes, the scheduled scripts' source and a bundled library, which must never be downloadable, and "
        "PHP files that WordPress loads from disk and that should never be requested directly. "
        "On Apache, the bundled `.htaccess` blocks them. Make sure `AllowOverride` permits it for the plugins directory. "
        "On nginx, add rules like these inside the server block:")
    g.code("""location ~* /wp-content/plugins/[^/]+/(cron|docs|tests|includes/lib)/ { return 404; }
location ~* /wp-content/plugins/[^/]+/.*\\.(md|txt|inc|log|sh|json|bak)$ { deny all; }
location ~* /wp-content/plugins/seal-ill-system/.*\\.(php|phtml|phar)$ { deny all; }""")
    g.p("Test it: requesting the plugin's README.md file, or any of its .php files, in a browser should give 403 or 404. Behind a CDN, "
        "clear its cache for the plugin folder after changing these rules, or an earlier answer may still be served.")
    g.h2("3.6 Configure, in this order")
    g.steps([
        "**Branding**: names, web address, logo, colors (section 4.2).",
        "**Contacts & Email**: alert addresses and From address. Turn **Email test mode** on while you set up. "
        "Change the wording of any email under **Email Subjects & Text** if you want to (section 4.3).",
        "**Search & Catalog**: the catalog addresses and search categories (section 4.5).",
        "**Users & Registration** and **Modules**: registration rules and the optional features you want (sections 4.6, 4.7).",
        "**Org Rules**: library systems, couriers, catalog matching and item types (section 4.8).",
        "Add your libraries under **Manage All Libraries** on the site (section 6.4).",
        "**Security**: login protection, password rules and two-factor authentication (section 4.10).",
        "Install the **scheduled scripts** (section 5) and paste the crontab from **Cron Jobs** (section 4.9).",
        "Create accounts for a few staff, run test requests end to end, then turn Email test mode off.",
    ])
