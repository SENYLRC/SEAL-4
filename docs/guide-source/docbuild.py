# Helpers for building the ILL System guide with python-docx.
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x1F, 0x2A, 0x44)      # headings
MUTED = RGBColor(0x50, 0x57, 0x5E)
ACCENT = "2F5597"                      # table header / borders
CODE_BG = "F2F2F2"
NOTE_BG = {"note": "EAF1FB", "warn": "FFF4CE", "tip": "E8F5E9"}
NOTE_BAR = {"note": "2F5597", "warn": "C59000", "tip": "2E7D32"}
NOTE_LABEL = {"note": "Note", "warn": "Important", "tip": "Tip"}


class Guide:
    def __init__(self):
        self.doc = Document()
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(1))
        sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.9)

        st = d.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for lvl, size, before in ((1, 18, 18), (2, 14, 14), (3, 12, 10)):
            h = st[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = INK
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            rpr = h.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                fonts.set(qn(a), "Calibri")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if fonts.get(qn(a)) is not None:
                    del fonts.attrib[qn(a)]
        self._num = 0

    # ── inline text: **bold**, `code`, *italic* ─────────────────────────────
    def _runs(self, par, text, size=None, color=None, bold=None):
        parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = par.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(9.5)
                r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = par.add_run(part[1:-1]); r.italic = True
            else:
                r = par.add_run(part)
            if size and not (part.startswith("`")):
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
            if bold is not None and not part.startswith("**"):
                r.bold = bold
        return par

    def h1(self, t, new_page=True):
        if new_page:
            self.page_break()
        return self.doc.add_heading(t, level=1)

    def h2(self, t):
        return self.doc.add_heading(t, level=2)

    def h3(self, t):
        return self.doc.add_heading(t, level=3)

    def p(self, t, **kw):
        par = self.doc.add_paragraph()
        self._runs(par, t, **kw)
        return par

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def bullets(self, items, level=0):
        for it in items:
            if isinstance(it, (list, tuple)):
                self.bullets(it, level + 1)
                continue
            par = self.doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
            par.paragraph_format.space_after = Pt(3)
            self._runs(par, it)

    def steps(self, items):
        """Numbered steps that always start at 1 (hand-numbered, hanging indent)."""
        n = 0
        for it in items:
            if isinstance(it, (list, tuple)):
                for sub in it:
                    par = self.doc.add_paragraph()
                    pf = par.paragraph_format
                    pf.left_indent = Inches(0.6); pf.first_line_indent = Inches(-0.2); pf.space_after = Pt(2)
                    self._runs(par, "– " + sub)
                continue
            n += 1
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Inches(0.35); pf.first_line_indent = Inches(-0.3); pf.space_after = Pt(4)
            pf.tab_stops.add_tab_stop(Inches(0.35))
            r = par.add_run(f"{n}.\t"); r.bold = True
            self._runs(par, it)

    def code(self, text):
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
            pf.left_indent = Inches(0.15); pf.keep_with_next = len(lines) <= 15 and i < len(lines) - 1   # short blocks stay on one page
            if i == len(lines) - 1:
                pf.space_after = Pt(8)
            self._shade(par, CODE_BG)
            r = par.add_run(line if line else " ")
            r.font.name = "Courier New"; r.font.size = Pt(8.5)
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")

    def note(self, text, kind="note"):
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.left_indent = Inches(0.1); pf.right_indent = Inches(0.1)
        pf.space_before = Pt(4); pf.space_after = Pt(8)
        self._shade(par, NOTE_BG[kind])
        self._border(par, "left", NOTE_BAR[kind], 24)
        r = par.add_run(NOTE_LABEL[kind] + ": "); r.bold = True
        self._runs(par, text)

    def table(self, headers, rows, widths):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        self._tbl_borders(t)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.width = Inches(widths[i])
            par = c.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            r = par.add_run(h); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.size = Pt(9.5)
            self._cell_shade(c, ACCENT)
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                c = cells[i]
                c.width = Inches(widths[i])
                chunks = val if isinstance(val, (list, tuple)) else [val]
                for j, chunk in enumerate(chunks):
                    par = c.paragraphs[0] if j == 0 else c.add_paragraph()
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.line_spacing = 1.0
                    self._runs(par, chunk, size=9.5)
        # column widths at grid level too (Word honours tblGrid)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i] * 1440)))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def toc(self):
        """A Word TOC field. Its visible result is filled in by save() with the heading list, so the contents show
        even before Word refreshes the field (Word then adds page numbers when it updates fields on open)."""
        self._toc_par = self.doc.add_paragraph()
        # ask Word to refresh fields (TOC, page numbers) when the file is opened
        settings = self.doc.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        settings.append(uf)

    def footer(self, text):
        sec = self.doc.sections[0]
        par = sec.footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(text + "   |   Page "); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
        self._field(par, "PAGE")

    # ── low-level XML helpers ───────────────────────────────────────────────
    def _field(self, par, code):
        r = par.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        ins = OxmlElement("w:instrText"); ins.set(qn("xml:space"), "preserve"); ins.text = code
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        r._r.append(b); r._r.append(ins); r._r.append(e)

    def _shade(self, par, fill):
        ppr = par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        ppr.append(shd)

    def _border(self, par, side, color, size):
        ppr = par._p.get_or_add_pPr()
        bdr = ppr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr"); ppr.append(bdr)
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:space"), "6"); el.set(qn("w:color"), color)
        bdr.append(el)

    def _cell_shade(self, cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    def _repeat_header(self, row):
        trpr = row._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trpr.append(h)

    def _tbl_borders(self, t):
        tblpr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "B4BCC8")
            b.append(el)
        tblpr.append(b)
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 50), ("bottom", 50), ("left", 90), ("right", 90)):
            el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(w)); el.set(qn("w:type"), "dxa"); mar.append(el)
        tblpr.append(mar)

    def _fill_toc(self):
        import copy
        heads = [(int(p.style.name[-1]), p.text) for p in self.doc.paragraphs
                 if p.style.name in ("Heading 1", "Heading 2") and p.text != "Contents"]
        anchor = self._toc_par
        def fld(run, kind):
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
        first = True
        prev = anchor
        for i, (lvl, text) in enumerate(heads):
            par = anchor if first else None
            if par is None:
                new_p = OxmlElement("w:p"); prev._p.addnext(new_p)
                from docx.text.paragraph import Paragraph
                par = Paragraph(new_p, anchor._parent)
            pf = par.paragraph_format
            pf.left_indent = Inches(0 if lvl == 1 else 0.3); pf.space_after = Pt(1 if lvl == 2 else 3)
            pf.space_before = Pt(6 if lvl == 1 else 0)
            if first:
                r0 = par.add_run(); fld(r0, "begin")
                ins = OxmlElement("w:instrText"); ins.set(qn("xml:space"), "preserve"); ins.text = 'TOC \\o "1-2" \\h \\z \\u'
                r0._r.append(ins); fld(r0, "separate")
            r = par.add_run(text); r.bold = (lvl == 1); r.font.size = Pt(10.5 if lvl == 1 else 10)
            if i == len(heads) - 1:
                fld(par.add_run(), "end")
            first = False
            prev = par
        settings = self.doc.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        settings.append(uf)

    def save(self, path):
        if getattr(self, "_toc_par", None) is not None:
            self._fill_toc()
        self.doc.save(path)
